import calendar
import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import Decimal
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from backend.models.financial_transaction import FinancialTransaction


ExportFormat = Literal["csv", "xlsx"]
TransactionType = Literal["expense", "income"]

CSV_MEDIA_TYPE = "text/csv; charset=utf-8"
XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

HEADERS = ("Data", "Descrição", "Categoria", "Tipo", "Valor (R$)", "Origem")
MONTH_NAMES = {
    "janeiro": 1,
    "fevereiro": 2,
    "marco": 3,
    "abril": 4,
    "maio": 5,
    "junho": 6,
    "julho": 7,
    "agosto": 8,
    "setembro": 9,
    "outubro": 10,
    "novembro": 11,
    "dezembro": 12,
}
SOURCE_LABELS = {
    "whatsapp_text": "WhatsApp texto",
    "whatsapp_audio": "WhatsApp áudio",
    "whatsapp_image": "WhatsApp imagem",
    "whatsapp_document": "WhatsApp documento",
    "dashboard_manual": "Dashboard manual",
    "web": "Web",
    "import": "Importação",
    "manual": "Manual",
}


@dataclass(frozen=True)
class StatementRequest:
    start_date: date
    end_date: date
    label: str
    transaction_type: TransactionType | None = None


@dataclass(frozen=True)
class ExportedStatement:
    content: bytes
    media_type: str
    filename: str
    transaction_count: int
    period_label: str


def current_month_period(current_date: date) -> StatementRequest:
    return StatementRequest(
        start_date=current_date.replace(day=1),
        end_date=current_date,
        label=current_date.strftime("%m/%Y"),
    )


def custom_period(start_date: date, end_date: date) -> StatementRequest:
    if start_date > end_date:
        raise ValueError("A data inicial deve ser anterior ou igual à data final")
    return StatementRequest(
        start_date=start_date,
        end_date=end_date,
        label=f"{start_date.strftime('%d/%m/%Y')} a {end_date.strftime('%d/%m/%Y')}",
    )


def resolve_whatsapp_statement_request(
    text: str,
    *,
    current_date: date,
) -> StatementRequest | None:
    normalized = _normalize_text(text)
    if not _is_statement_request(normalized):
        return None

    transaction_type: TransactionType | None = None
    if re.search(r"\b(gasto|gastos|despesa|despesas)\b", normalized):
        transaction_type = "expense"
    elif re.search(r"\b(receita|receitas|entrada|entradas)\b", normalized):
        transaction_type = "income"

    if re.search(r"\bultimos?\s+30\s+dias\b", normalized):
        return StatementRequest(
            start_date=current_date - timedelta(days=29),
            end_date=current_date,
            label="últimos 30 dias",
            transaction_type=transaction_type,
        )

    if re.search(r"\b(dessa|desta|essa|nesta)\s+semana\b", normalized):
        return StatementRequest(
            start_date=current_date - timedelta(days=current_date.weekday()),
            end_date=current_date,
            label="esta semana",
            transaction_type=transaction_type,
        )

    if re.search(r"\bmes\s+passado\b|\bdo\s+mes\s+anterior\b", normalized):
        current_month_start = current_date.replace(day=1)
        previous_month_end = current_month_start - timedelta(days=1)
        return StatementRequest(
            start_date=previous_month_end.replace(day=1),
            end_date=previous_month_end,
            label=previous_month_end.strftime("%m/%Y"),
            transaction_type=transaction_type,
        )

    month_match = re.search(
        rf"\b({'|'.join(MONTH_NAMES)})\b(?:\s+de\s+(\d{{4}}))?",
        normalized,
    )
    if month_match:
        month = MONTH_NAMES[month_match.group(1)]
        year = int(month_match.group(2)) if month_match.group(2) else current_date.year
        if month_match.group(2) is None and month > current_date.month:
            year -= 1
        last_day = calendar.monthrange(year, month)[1]
        return StatementRequest(
            start_date=date(year, month, 1),
            end_date=date(year, month, last_day),
            label=f"{month:02d}/{year}",
            transaction_type=transaction_type,
        )

    if re.search(
        r"\b(do|desse|deste|esse|neste)\s+mes\b|\bmes\s+atual\b",
        normalized,
    ):
        period = current_month_period(current_date)
        return StatementRequest(
            start_date=period.start_date,
            end_date=period.end_date,
            label=period.label,
            transaction_type=transaction_type,
        )

    return None


def generate_statement(
    db: Session,
    *,
    user_id: int,
    period: StatementRequest,
    export_format: ExportFormat = "xlsx",
) -> ExportedStatement | None:
    transactions = _list_transactions(
        db,
        user_id=user_id,
        period=period,
    )
    if not transactions:
        return None

    period_slug = _period_filename_slug(period)
    if export_format == "csv":
        content = _build_csv(transactions)
        media_type = CSV_MEDIA_TYPE
    elif export_format == "xlsx":
        content = _build_xlsx(transactions)
        media_type = XLSX_MEDIA_TYPE
    else:
        raise ValueError("Formato de exportação inválido")

    return ExportedStatement(
        content=content,
        media_type=media_type,
        filename=f"extrato-fincontrol-{period_slug}.{export_format}",
        transaction_count=len(transactions),
        period_label=period.label,
    )


def _list_transactions(
    db: Session,
    *,
    user_id: int,
    period: StatementRequest,
) -> list[FinancialTransaction]:
    statement = (
        select(FinancialTransaction)
        .options(joinedload(FinancialTransaction.category))
        .where(
            FinancialTransaction.user_id == user_id,
            FinancialTransaction.transaction_date >= period.start_date,
            FinancialTransaction.transaction_date <= period.end_date,
        )
        .order_by(
            FinancialTransaction.transaction_date.asc(),
            FinancialTransaction.id.asc(),
        )
    )
    if period.transaction_type is not None:
        statement = statement.where(
            FinancialTransaction.type == period.transaction_type
        )
    return list(db.scalars(statement))


def _build_csv(transactions: list[FinancialTransaction]) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter=";", lineterminator="\r\n")
    writer.writerow(HEADERS)
    for transaction in transactions:
        writer.writerow(
            (
                transaction.transaction_date.strftime("%d/%m/%Y"),
                _safe_spreadsheet_text(transaction.description),
                _safe_spreadsheet_text(_category_name(transaction)),
                _type_label(transaction.type),
                _format_decimal_pt_br(transaction.amount),
                _source_label(transaction.source),
            )
        )
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def _build_xlsx(transactions: list[FinancialTransaction]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Extrato"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"

    worksheet.append(HEADERS)
    for transaction in transactions:
        worksheet.append(
            (
                transaction.transaction_date,
                _safe_spreadsheet_text(transaction.description),
                _safe_spreadsheet_text(_category_name(transaction)),
                _type_label(transaction.type),
                float(Decimal(transaction.amount)),
                _source_label(transaction.source),
            )
        )

    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 24

    for row in worksheet.iter_rows(min_row=2):
        row[0].number_format = "DD/MM/YYYY"
        row[0].alignment = Alignment(horizontal="center")
        row[4].number_format = 'R$ #,##0.00;[Red]-R$ #,##0.00'
        row[4].alignment = Alignment(horizontal="right")
        for cell in (row[1], row[2], row[3], row[5]):
            cell.alignment = Alignment(vertical="center")

    column_widths = (12, 40, 22, 12, 16, 22)
    for index, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    worksheet.auto_filter.ref = f"A1:F{worksheet.max_row}"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _is_statement_request(normalized: str) -> bool:
    period_pattern = (
        r"\b(mes|semana|dias|janeiro|fevereiro|marco|abril|maio|junho|"
        r"julho|agosto|setembro|outubro|novembro|dezembro)\b"
    )
    has_period = re.search(period_pattern, normalized) is not None
    has_statement = re.search(r"\bextrato\b", normalized) is not None
    has_movements = re.search(r"\bmovimentacoes\b", normalized) is not None
    has_expenses = re.search(r"\b(gasto|gastos|despesa|despesas)\b", normalized) is not None
    asks_to_send = re.search(r"\b(manda|mandar|envia|enviar|gera|gerar)\b", normalized) is not None
    return has_period and (
        has_statement
        or has_movements
        or (has_expenses and asks_to_send)
    )


def _normalize_text(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value.casefold().strip())
    without_accents = "".join(
        character
        for character in normalized
        if not unicodedata.combining(character)
    )
    return " ".join(
        "".join(
            character if character.isalnum() else " "
            for character in without_accents
        ).split()
    )


def _period_filename_slug(period: StatementRequest) -> str:
    if (
        period.start_date.year == period.end_date.year
        and period.start_date.month == period.end_date.month
    ):
        return period.start_date.strftime("%Y-%m")
    return f"{period.start_date.isoformat()}-a-{period.end_date.isoformat()}"


def _category_name(transaction: FinancialTransaction) -> str:
    return transaction.category.name if transaction.category is not None else "Sem categoria"


def _type_label(transaction_type: str) -> str:
    return "Entrada" if transaction_type == "income" else "Saída"


def _source_label(source: str) -> str:
    return SOURCE_LABELS.get(source, source.replace("_", " ").title())


def _format_decimal_pt_br(value: Decimal) -> str:
    return f"{Decimal(value):.2f}".replace(".", ",")


def _safe_spreadsheet_text(value: str) -> str:
    normalized = " ".join(value.split())
    if normalized.startswith(("=", "+", "-", "@")):
        return f"'{normalized}"
    return normalized
