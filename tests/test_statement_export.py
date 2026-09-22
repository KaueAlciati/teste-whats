import asyncio
import csv
import io
import os
import unittest
from datetime import date, datetime
from decimal import Decimal
from unittest.mock import AsyncMock, patch
from zoneinfo import ZoneInfo

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.database.base import Base
from backend.database.connection import get_db
from backend.main import app
from backend.models import Category, FinancialTransaction, User
from backend.services.auth_service import create_access_token
from backend.services.financial_assistant_service import (
    handle_financial_message,
    process_financial_message,
)
from backend.services.statement_export_service import (
    ExportedStatement,
    current_month_period,
    custom_period,
    generate_statement,
    resolve_whatsapp_statement_request,
)


class StatementExportTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.engine = create_engine(
            "sqlite+pysqlite:///:memory:",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.session_factory = sessionmaker(
            bind=self.engine,
            expire_on_commit=False,
        )

        def override_get_db():
            with self.session_factory() as db:
                yield db

        app.dependency_overrides[get_db] = override_get_db
        self.environment = patch.dict(
            os.environ,
            {"JWT_SECRET_KEY": "test-secret-with-at-least-thirty-two-characters"},
        )
        self.environment.start()
        self.client = TestClient(app)

        with self.session_factory() as db:
            self.category = Category(
                name="Alimentação",
                type="expense",
                is_default=True,
            )
            self.user = self._user(db, "user@example.com", "5515999999999")
            self.other_user = self._user(db, "other@example.com", "5515888888888")
            db.add(self.category)
            db.commit()
            db.refresh(self.category)
            self.token = create_access_token(self.user)

    def tearDown(self) -> None:
        self.client.close()
        app.dependency_overrides.clear()
        self.environment.stop()
        Base.metadata.drop_all(self.engine)
        self.engine.dispose()

    def test_csv_export_has_expected_columns_and_isolates_user(self) -> None:
        with self.session_factory() as db:
            self._transaction(db, self.user.id, date(2026, 9, 5), "Mercado")
            self._transaction(db, self.other_user.id, date(2026, 9, 5), "Privada")
            exported = generate_statement(
                db,
                user_id=self.user.id,
                period=custom_period(date(2026, 9, 1), date(2026, 9, 30)),
                export_format="csv",
            )

        self.assertIsNotNone(exported)
        decoded = exported.content.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(decoded), delimiter=";"))
        self.assertEqual(
            rows[0],
            ["Data", "Descrição", "Categoria", "Tipo", "Valor (R$)", "Origem"],
        )
        self.assertEqual(rows[1], [
            "05/09/2026",
            "Mercado",
            "Alimentação",
            "Saída",
            "40,50",
            "Dashboard manual",
        ])
        self.assertNotIn("Privada", decoded)

    def test_xlsx_export_contains_real_types_and_formatting(self) -> None:
        with self.session_factory() as db:
            self._transaction(db, self.user.id, date(2026, 9, 7), "Padaria")
            exported = generate_statement(
                db,
                user_id=self.user.id,
                period=custom_period(date(2026, 9, 1), date(2026, 9, 30)),
                export_format="xlsx",
            )

        workbook = load_workbook(io.BytesIO(exported.content))
        worksheet = workbook["Extrato"]
        self.assertEqual([cell.value for cell in worksheet[1]], list(
            ("Data", "Descrição", "Categoria", "Tipo", "Valor (R$)", "Origem")
        ))
        exported_date = worksheet["A2"].value
        if isinstance(exported_date, datetime):
            exported_date = exported_date.date()
        self.assertEqual(exported_date, date(2026, 9, 7))
        self.assertEqual(worksheet["A2"].number_format, "DD/MM/YYYY")
        self.assertEqual(worksheet["B2"].value, "Padaria")
        self.assertEqual(worksheet["E2"].value, 40.5)
        self.assertIn("R$", worksheet["E2"].number_format)
        self.assertEqual(worksheet["A1"].fill.fgColor.rgb, "000F766E")
        self.assertTrue(worksheet["A1"].font.bold)
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertGreaterEqual(worksheet.column_dimensions["B"].width, 35)
        workbook.close()

    def test_api_defaults_to_current_month_and_ignores_user_id(self) -> None:
        today = datetime.now(ZoneInfo("America/Sao_Paulo")).date()
        with self.session_factory() as db:
            self._transaction(db, self.user.id, today, "Atual")
            self._transaction(db, self.other_user.id, today, "Outro usuário")

        response = self.client.get(
            "/api/transactions/export",
            params={"format": "xlsx", "user_id": self.other_user.id},
            headers={**self._headers(), "Origin": "http://localhost:3000"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "Content-Disposition",
            response.headers["access-control-expose-headers"],
        )
        workbook = load_workbook(io.BytesIO(response.content))
        descriptions = [cell.value for cell in workbook["Extrato"]["B"]]
        self.assertIn("Atual", descriptions)
        self.assertNotIn("Outro usuário", descriptions)
        workbook.close()

    def test_api_exports_custom_interval_as_csv(self) -> None:
        with self.session_factory() as db:
            self._transaction(db, self.user.id, date(2026, 8, 10), "Dentro")
            self._transaction(db, self.user.id, date(2026, 8, 25), "Fora")

        response = self.client.get(
            "/api/transactions/export",
            params={
                "format": "csv",
                "start_date": "2026-08-01",
                "end_date": "2026-08-15",
            },
            headers=self._headers(),
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))
        text = response.content.decode("utf-8-sig")
        self.assertIn("Dentro", text)
        self.assertNotIn("Fora", text)
        self.assertIn(".csv", response.headers["content-disposition"])

    def test_specific_month_and_supported_natural_phrases(self) -> None:
        expected = {
            "me manda meu extrato do mês": (date(2026, 9, 1), date(2026, 9, 21), None),
            "extrato desse mês": (date(2026, 9, 1), date(2026, 9, 21), None),
            "gera meu extrato de setembro": (date(2026, 9, 1), date(2026, 9, 30), None),
            "me manda os gastos dessa semana": (date(2026, 9, 21), date(2026, 9, 21), "expense"),
            "extrato dos últimos 30 dias": (date(2026, 8, 23), date(2026, 9, 21), None),
            "me manda minhas movimentações de agosto": (date(2026, 8, 1), date(2026, 8, 31), None),
        }
        for phrase, bounds in expected.items():
            with self.subTest(phrase=phrase):
                resolved = resolve_whatsapp_statement_request(
                    phrase,
                    current_date=date(2026, 9, 21),
                )
                self.assertIsNotNone(resolved)
                self.assertEqual(
                    (resolved.start_date, resolved.end_date, resolved.transaction_type),
                    bounds,
                )

    def test_no_result_does_not_generate_empty_file(self) -> None:
        with self.session_factory() as db:
            exported = generate_statement(
                db,
                user_id=self.user.id,
                period=current_month_period(date(2026, 9, 21)),
                export_format="xlsx",
            )
        self.assertIsNone(exported)

        response = self.client.get(
            "/api/transactions/export",
            headers=self._headers(),
        )
        self.assertEqual(response.status_code, 404)

    def test_whatsapp_text_and_audio_use_same_export_resolver(self) -> None:
        with self.session_factory() as db:
            self._transaction(db, self.user.id, date(2026, 9, 10), "WhatsApp")
            text_result = handle_financial_message(
                db,
                user=db.get(User, self.user.id),
                text="me manda meu extrato do mês",
                whatsapp_message_id="wamid.statement-text",
                current_date=date(2026, 9, 21),
            )
            audio_result = handle_financial_message(
                db,
                user=db.get(User, self.user.id),
                text="extrato desse mês",
                whatsapp_message_id="wamid.statement-audio",
                current_date=date(2026, 9, 21),
                source="whatsapp_audio",
                audio_transcription="extrato desse mês",
            )

        self.assertIsInstance(text_result, ExportedStatement)
        self.assertIsInstance(audio_result, ExportedStatement)
        text_workbook = load_workbook(io.BytesIO(text_result.content))
        audio_workbook = load_workbook(io.BytesIO(audio_result.content))
        self.assertEqual(
            text_workbook["Extrato"]["B2"].value,
            audio_workbook["Extrato"]["B2"].value,
        )
        text_workbook.close()
        audio_workbook.close()
        self.assertTrue(text_result.filename.endswith("2026-09.xlsx"))

    def test_whatsapp_delivery_uses_document_sender(self) -> None:
        exported = ExportedStatement(
            content=b"xlsx-bytes",
            media_type="application/test-xlsx",
            filename="extrato-fincontrol-2026-09.xlsx",
            transaction_count=1,
            period_label="09/2026",
        )
        document_mock = AsyncMock(return_value=True)
        text_mock = AsyncMock(return_value=True)
        with (
            patch(
                "backend.services.financial_assistant_service._process_financial_message",
                return_value=exported,
            ),
            patch(
                "backend.services.financial_assistant_service.send_document_message",
                document_mock,
            ),
            patch(
                "backend.services.financial_assistant_service.send_text_message",
                text_mock,
            ),
        ):
            asyncio.run(
                process_financial_message(
                    "5515999999999",
                    "wamid.delivery",
                    "extrato desse mês",
                )
            )

        document_mock.assert_awaited_once_with(
            "5515999999999",
            content=b"xlsx-bytes",
            filename="extrato-fincontrol-2026-09.xlsx",
            mime_type="application/test-xlsx",
            caption="Extrato FinControl AI · 09/2026",
        )
        text_mock.assert_not_awaited()

    def _transaction(
        self,
        db,
        user_id: int,
        transaction_date: date,
        description: str,
    ) -> FinancialTransaction:
        transaction = FinancialTransaction(
            user_id=user_id,
            type="expense",
            amount=Decimal("40.50"),
            description=description,
            category_id=self.category.id,
            transaction_date=transaction_date,
            source="dashboard_manual",
        )
        db.add(transaction)
        db.commit()
        db.refresh(transaction)
        return transaction

    def _headers(self) -> dict[str, str]:
        return {"Authorization": f"Bearer {self.token}"}

    @staticmethod
    def _user(db, email: str, phone: str) -> User:
        user = User(
            name="Usuário",
            email=email,
            whatsapp_phone=phone,
            password_hash="$argon2id$test-hash",
        )
        db.add(user)
        db.commit()
        db.refresh(user)
        return user


if __name__ == "__main__":
    unittest.main()
